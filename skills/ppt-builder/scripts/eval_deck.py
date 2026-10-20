# -*- coding: utf-8 -*-
"""对 ppt-builder 产出的 pptx 下客观断言——技能改版/回归时不用每次现写探针。

只做几何与存在性判断（不做审美）：页数、讲者备注覆盖率、整页图的铺满程度、
有没有在整页图上叠多余文字，以及复用同目录 qa.py 的几何核对结果。

用法：
    python eval_deck.py 演示.pptx                        # 常规内容稿（只查页数/备注/qa）
    python eval_deck.py 还原.pptx --mode restore --pages 10   # 整页图还原稿
    python eval_deck.py 演示.pptx --pages 6 --notes-min 60 --json

退出码：0 全部通过；1 有断言失败；2 文件打不开或不是 pptx。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
MAX_EXTRA_TEXT = 3          # 整页图页上允许出现的文字长度上限：页码这类短标记
TEXT_SHAPE = 17             # python-pptx MSO_SHAPE_TYPE.TEXT_BOX，属装饰性判断之外的文字形状


def _notes_len(slide) -> int:
    if not slide.has_notes_slide:
        return 0
    return len("".join(slide.notes_slide.notes_text_frame.text.split()))


def _picture_cover(prs, slide) -> float:
    """本页图片覆盖画布的面积比（多张取最大）。"""
    sw, sh = prs.slide_width, prs.slide_height
    best = 0.0
    for shp in slide.shapes:
        if shp.shape_type == 13:        # PICTURE
            best = max(best, (shp.width or 0) * (shp.height or 0) / (sw * sh))
    return best


def _long_texts(slide) -> list[str]:
    """页上长度超过 MAX_EXTRA_TEXT 的文字（整页图页应为空；页码/图注短标记不计）。"""
    out = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            t = shp.text_frame.text.strip()
            if len(t) > MAX_EXTRA_TEXT:
                out.append(t)
        if getattr(shp, "has_table", False) and shp.has_table:
            for row in shp.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if len(t) > MAX_EXTRA_TEXT:
                        out.append(t)
    return out


def _qa_counts(path: Path) -> tuple[int | None, int | None]:
    """复用同目录 qa.py 的几何核对：返回 (fail 数, warn 数)；不可用时返回 (None, None)。"""
    try:
        from qa import check_presentation      # 同目录脚本
    except Exception:
        return None, None
    try:
        res = check_presentation(str(path))
    except Exception:
        return None, None
    return res["counts"]["fail"], res["counts"]["warn"]


def evaluate(path: str, mode: str = "normal", pages: int | None = None,
             notes_min: int = 20, cover_min: float = 0.80) -> dict:
    from pptx import Presentation

    p = Path(path)
    prs = Presentation(str(p))
    n = len(prs.slides)
    checks: list[dict] = []

    def add(name: str, ok: bool, detail: str) -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    if pages is None:
        add("页数", True, f"{n} 页（未指定期望值，仅报告）")
    else:
        add("页数", n == pages, f"实际 {n} 页，期望 {pages} 页")

    short = [(i, _notes_len(s)) for i, s in enumerate(prs.slides, 1) if _notes_len(s) < notes_min]
    add("讲者备注", not short,
        "全部页备注齐全" if not short else f"备注不足 {notes_min} 字的页：{short}")

    if mode == "restore":
        covers = [(i, round(_picture_cover(prs, s), 2)) for i, s in enumerate(prs.slides, 1)]
        low = [(i, c) for i, c in covers if c < cover_min]
        add("整页图铺满", not low, f"最小覆盖比 {min((c for _, c in covers), default=0):.2f}"
            + ("" if not low else f"；未达标页：{low}"))

        extra = [(i, t[:18]) for i, s in enumerate(prs.slides, 1) for t in _long_texts(s)]
        add("图上无多余文字", not extra, "无叠加文字" if not extra else f"多余文字：{extra[:3]}")

    fails, warns = _qa_counts(p)
    if fails is None:
        add("qa 几何核对", True, "qa.py 不可用，已跳过")
    else:
        add("qa 几何核对", fails == 0, f"qa.py：fail {fails} / warn {warns}")

    failed = [c for c in checks if not c["ok"]]
    return {"file": str(p), "mode": mode, "slide_count": n,
            "passed": len(checks) - len(failed), "failed": len(failed), "checks": checks}


def main() -> int:
    ap = argparse.ArgumentParser(description="对 pptx 下客观断言（技能回归用）")
    ap.add_argument("pptx")
    ap.add_argument("--mode", choices=["normal", "restore"], default="normal")
    ap.add_argument("--pages", type=int, default=None, help="期望页数")
    ap.add_argument("--notes-min", type=int, default=20, help="讲者备注最少字数")
    ap.add_argument("--cover-min", type=float, default=0.80, help="整页图最小覆盖比（restore 模式）")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    try:
        res = evaluate(a.pptx, a.mode, a.pages, a.notes_min, a.cover_min)
    except Exception as e:                      # 打不开/不是 pptx
        print(f"无法读取：{a.pptx}（{type(e).__name__}: {e}）", file=sys.stderr)
        return 2

    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        print(f"{res['file']}  mode={res['mode']}  {res['slide_count']} 页  "
              f"通过 {res['passed']}/{res['passed'] + res['failed']}")
        for c in res["checks"]:
            print(f"  [{'PASS' if c['ok'] else 'FAIL'}] {c['name']}：{c['detail']}")
    return 0 if res["failed"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
