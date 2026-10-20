#!/usr/bin/env python3
"""qa.py 自测：纯标准库 + python-pptx，直接 ``python test_qa.py`` 运行（不依赖 pytest）。

覆盖：
1. 正常生成的 deck（PPT_QA_SAMPLE_DECK 指定的真实 deck，未设置时用 examples/demo.py 现场生成）→ fail 数必须为 0；
2. 人为构造形状越界 → 必须 fail；
3. 无讲者备注 → 必须 warn，且报告缺失页码；
4. 文字里带 TODO → 必须 fail；
5. 空文本框与纯装饰色块的分界口径；
6. 图片分辨率（低 DPI → fail，高 DPI → 不报）；
7. 标题重复（完全相同 → fail；仅标点 / 空白不同 → warn）；
8. 组合形状坐标换算（子形状按 a:xfrm 换算到幻灯片坐标）；
9. 字体清单（非 Microsoft YaHei 一并列出）；
10. CLI 契约（退出码 0 / 1 / 2，--json、--json-out）。
"""

from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import traceback
import zlib
from typing import Any, Callable, Dict, List, Optional

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
SKILL_DIR = os.path.dirname(SCRIPTS)
sys.path.insert(0, SCRIPTS)

from qa import (  # noqa: E402  （必须在 sys.path 调整之后导入）
    CHECK_BOUNDS, CHECK_EMPTY_TEXT, CHECK_FONTS, CHECK_IMAGE_DPI, CHECK_NOTES,
    CHECK_PLACEHOLDER, CHECK_TITLE, check_presentation,
)

BRAND_DECK = os.environ.get("PPT_QA_SAMPLE_DECK", "")
DEMO_PY = os.path.join(SKILL_DIR, "examples", "demo.py")
PYTHON = sys.executable

_TMP_DIR: Optional[str] = None
_DEMO_DECK: Optional[str] = None


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------
def tmp_dir() -> str:
    """本次自测的临时目录。"""
    global _TMP_DIR
    if _TMP_DIR is None:
        _TMP_DIR = tempfile.mkdtemp(prefix="qa-selftest-", dir="/tmp")
    return _TMP_DIR


def new_deck() -> Presentation:
    """960×540pt 的空 deck。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    return prs


def add_slide(prs: Presentation):
    """加一页空白页。"""
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(slide: Any, x: float, y: float, w: float, h: float, text: str,
             size: int = 24, name: Optional[str] = None):
    """加一个文本框（直接建形状，测试里允许）。"""
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    box.text_frame.text = text
    for para in box.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)
    if name:
        box.name = name
    return box


def add_panel(slide: Any, x: float, y: float, w: float, h: float, name: Optional[str] = None):
    """加一个纯色装饰色块（无文字、无 run，与 pptdeck 的装饰一致）。"""
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.line.fill.background()
    if name:
        shape.name = name
    return shape


def save(prs: Presentation, filename: str) -> str:
    """把 deck 存到临时目录并返回路径。"""
    path = os.path.join(tmp_dir(), filename)
    prs.save(path)
    return path


def png_bytes(width: int, height: int, rgb=(120, 140, 160)) -> bytes:
    """用标准库造一张纯色 PNG（避免测试依赖 Pillow）。"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def png_file(name: str, width: int, height: int) -> str:
    """写一张 PNG 到临时目录并返回路径。"""
    path = os.path.join(tmp_dir(), name)
    with open(path, "wb") as handle:
        handle.write(png_bytes(width, height))
    return path


def findings(result: Dict[str, Any], level: Optional[str] = None,
             check: Optional[str] = None) -> List[Dict[str, Any]]:
    """按级别 / 检查项过滤结论。"""
    return [f for f in result["findings"]
            if (level is None or f["level"] == level) and (check is None or f["check"] == check)]


def pages_of(result: Dict[str, Any], level: str, check: str) -> List[int]:
    """某项检查命中的页码列表。"""
    return sorted(f["page"] for f in findings(result, level, check) if f["page"])


def set_group_xfrm(group: Any, off: tuple, ext: tuple, ch_off: tuple, ch_ext: tuple) -> None:
    """直接改写组合形状的 a:xfrm（绕过 python-pptx 的自动重算），单位 pt。"""
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    for tag, (x, y) in (("a:off", off), ("a:chOff", ch_off)):
        el = xfrm.find(qn(tag))
        el.set("x", str(Pt(x).emu))
        el.set("y", str(Pt(y).emu))
    for tag, (cx, cy) in (("a:ext", ext), ("a:chExt", ch_ext)):
        el = xfrm.find(qn(tag))
        el.set("cx", str(Pt(cx).emu))
        el.set("cy", str(Pt(cy).emu))


def demo_deck() -> str:
    """正常 deck：优先用 ``PPT_QA_SAMPLE_DECK`` 指定的真实 deck；没有就用 examples/demo.py 现场生成。"""
    global _DEMO_DECK
    if BRAND_DECK and os.path.isfile(BRAND_DECK):
        return BRAND_DECK
    if _DEMO_DECK and os.path.isfile(_DEMO_DECK):
        return _DEMO_DECK
    if not os.path.isfile(DEMO_PY):
        raise AssertionError(
            f"既没有 PPT_QA_SAMPLE_DECK（当前 {BRAND_DECK or '未设置'}），也找不到 {DEMO_PY}")
    work = os.path.join(tmp_dir(), "demo")
    os.makedirs(work, exist_ok=True)
    source = open(DEMO_PY, encoding="utf-8").read().splitlines(keepends=True)
    patched = [f'sys.path.insert(0, r"{SCRIPTS}")\n'
               if "sys.path.insert" in line else line for line in source]
    with open(os.path.join(work, "demo.py"), "w", encoding="utf-8") as handle:
        handle.write("".join(patched))
    proc = subprocess.run([PYTHON, "demo.py"], cwd=work, capture_output=True, text=True)
    decks = [os.path.join(work, n) for n in sorted(os.listdir(work)) if n.endswith(".pptx")]
    if proc.returncode != 0 or not decks:
        raise AssertionError(
            "examples/demo.py 生成正常 deck 失败：\n"
            f"stdout={proc.stdout[-800:]}\nstderr={proc.stderr[-800:]}")
    _DEMO_DECK = decks[0]
    return _DEMO_DECK


# --------------------------------------------------------------------------
# 用例 1：正常 deck
# --------------------------------------------------------------------------
def test_normal_deck_has_no_fail() -> str:
    deck = demo_deck()
    captured = io.StringIO()
    with contextlib.redirect_stdout(captured), contextlib.redirect_stderr(captured):
        result = check_presentation(deck)
    assert captured.getvalue() == "", f"check_presentation 不能有副作用（打印了 {captured.getvalue()[:80]!r}）"
    assert result["slide_count"] > 0, "页数应大于 0"
    assert result["counts"]["fail"] == 0, f"正常 deck 不应有 fail：{findings(result, 'fail')[:3]}"
    assert not result["errors"], f"检查过程不应报错：{result['errors'][:3]}"
    assert result["fonts"], "应收集到字体清单"
    # warn 的零告警断言只在真实 deck 上成立：examples/demo.py 生成的示例 deck 自带
    # 判断题选项空文本框与库自动生成的答案页（无备注入口），这部分告警由用例 3、5 单独覆盖。
    if os.environ.get("PPT_QA_SAMPLE_DECK"):
        assert result["counts"]["warn"] == 0, f"正常 deck 不应有 warn：{findings(result, 'warn')[:3]}"
        return f"（{os.path.basename(deck)}，{result['slide_count']} 页，fail 0 / warn 0）"
    return f"（{os.path.basename(deck)}，{result['slide_count']} 页，fail 0；示例 deck 不校验 warn）"


# --------------------------------------------------------------------------
# 用例 2：形状越界
# --------------------------------------------------------------------------
def test_out_of_bounds() -> str:
    prs = new_deck()
    slide = add_slide(prs)
    add_panel(slide, 0, 0, 960, 540, "满版底")          # 贴边不算越界（容差 1pt）
    add_text(slide, 900, 200, 300, 60, "右边超出的文本框", 24, "越界框")
    result = check_presentation(save(prs, "bounds.pptx"))
    hits = findings(result, "fail", CHECK_BOUNDS)
    assert hits, "越界形状必须报 fail"
    assert any("越界框" in f["message"] for f in hits), f"应点名越界形状：{hits}"
    assert all(f["page"] == 1 for f in hits), "越界应记在第 1 页"
    assert not any("满版底" in f["message"] for f in hits), "贴边形状不应报越界"
    return f"（报出 {len(hits)} 项越界：{hits[0]['message'][:40]}…）"


# --------------------------------------------------------------------------
# 用例 3：讲者备注覆盖率
# --------------------------------------------------------------------------
def test_missing_notes() -> str:
    prs = new_deck()
    for index in range(1, 4):
        slide = add_slide(prs)
        add_text(slide, 48, 24, 864, 84, f"第{index}页的结论", 32)
        if index == 2:
            slide.notes_slide.notes_text_frame.text = "这页有足够长的讲者备注，讲者照着能讲下去。"
    result = check_presentation(save(prs, "notes.pptx"))
    hits = findings(result, "warn", CHECK_NOTES)
    assert hits, "缺备注必须报 warn"
    assert pages_of(result, "warn", CHECK_NOTES) == [1, 3], \
        f"应报出第 1、3 页：{pages_of(result, 'warn', CHECK_NOTES)}"
    # 备注字数下限可调：把下限降到 5 字，第 2 页不算短
    loose = check_presentation(save(prs, "notes2.pptx"), notes_min=5)
    assert 2 not in pages_of(loose, "warn", CHECK_NOTES), "备注达到下限就不该报"
    return f"（缺备注页码 {pages_of(result, 'warn', CHECK_NOTES)}）"


# --------------------------------------------------------------------------
# 用例 4：占位残留
# --------------------------------------------------------------------------
def test_placeholder_text() -> str:
    prs = new_deck()
    slide = add_slide(prs)
    add_text(slide, 48, 24, 864, 84, "TODO：把这段换成真数据", 32)
    add_text(slide, 48, 200, 600, 60, "这里先写 Lorem ipsum", 24)
    second = add_slide(prs)
    add_text(second, 48, 24, 864, 84, "正常标题", 32)
    second.notes_slide.notes_text_frame.text = "备注里还留着占位符，交付前要删掉。"
    result = check_presentation(save(prs, "todo.pptx"))
    hits = findings(result, "fail", CHECK_PLACEHOLDER)
    assert len(hits) >= 2, f"页面文字与备注里的占位词都要报：{hits}"
    assert any("TODO" in f["message"] for f in hits), "应点名 TODO"
    assert any(f["page"] == 2 for f in hits), "备注里的占位词应记在它所在页"
    assert result["counts"]["fail"] >= 2
    return f"（报出 {len(hits)} 项：{hits[0]['message'][:40]}…）"


# --------------------------------------------------------------------------
# 用例 5：空文本框与纯装饰色块
# --------------------------------------------------------------------------
def test_empty_textbox_versus_decoration() -> str:
    prs = new_deck()
    slide = add_slide(prs)
    add_panel(slide, 0, 0, 960, 8, "顶部色条")           # 8pt 色条：装饰
    add_panel(slide, 0, 12, 960, 360, "封面大色块")       # 大色块：装饰
    slide.shapes.add_textbox(Pt(48), Pt(140), Pt(600), Pt(200)).name = "忘了填的文本框"
    card = add_text(slide, 48, 380, 600, 120, "正文卡片", 24, "被清空的卡片")
    add_panel(slide, 48, 380, 600, 120, "卡片底板")
    card.text_frame.paragraphs[0].runs[0].text = ""       # 有 run 但文字为空：疑似丢字
    result = check_presentation(save(prs, "empty.pptx"))
    hits = findings(result, "warn", CHECK_EMPTY_TEXT)
    names = " ".join(f["message"] for f in hits)
    assert "忘了填的文本框" in names, f"空文本框必须报：{hits}"
    assert "被清空的卡片" in names, f"文字被清空的卡片必须报：{hits}"
    assert "顶部色条" not in names and "封面大色块" not in names and "卡片底板" not in names, \
        f"纯装饰色块不该报：{hits}"
    return f"（报出 {len(hits)} 项，装饰色块已排除）"


# --------------------------------------------------------------------------
# 用例 6：图片分辨率
# --------------------------------------------------------------------------
def test_image_dpi() -> str:
    prs = new_deck()
    slide = add_slide(prs)
    add_text(slide, 48, 24, 864, 84, "图片分辨率检查", 32)
    slide.shapes.add_picture(png_file("low.png", 60, 60), Pt(48), Pt(140), Pt(144), Pt(144))
    slide.shapes.add_picture(png_file("high.png", 600, 600), Pt(400), Pt(140), Pt(144), Pt(144))
    result = check_presentation(save(prs, "dpi.pptx"))
    hits = findings(result, level=None, check=CHECK_IMAGE_DPI)
    assert len(hits) == 1, f"只应报低分辨率那张：{hits}"
    assert hits[0]["level"] == "fail" and hits[0]["page"] == 1, f"30 DPI 必须 fail：{hits}"
    assert "30 DPI" in hits[0]["message"], f"应报出估算 DPI：{hits[0]['message']}"
    return f"（{hits[0]['message']}）"


# --------------------------------------------------------------------------
# 用例 7：标题重复
# --------------------------------------------------------------------------
def test_duplicate_titles() -> str:
    prs = new_deck()
    for text in ("同一句结论", "同一句结论", "结论B：先做诊断", "结论B 先做诊断"):
        add_text(add_slide(prs), 48, 24, 864, 84, text, 32)
    result = check_presentation(save(prs, "titles.pptx"))
    fails = findings(result, "fail", CHECK_TITLE)
    warns = findings(result, "warn", CHECK_TITLE)
    assert fails and fails[0]["page"] == 2, f"完全相同的标题必须 fail：{fails}"
    assert warns and warns[0]["page"] == 4, f"仅标点不同的标题应 warn：{warns}"
    return f"（fail 1 项 / warn 1 项：{warns[0]['message'][:46]}…）"


# --------------------------------------------------------------------------
# 用例 8：组合形状坐标换算
# --------------------------------------------------------------------------
def test_group_shape_transform() -> str:
    prs = new_deck()
    slide = add_slide(prs)

    inside = slide.shapes.add_group_shape()               # 缩放 0.5：子形状换算后落在画布内
    inner = inside.shapes.add_textbox(Pt(1800), Pt(0), Pt(100), Pt(50))
    inner.text_frame.text = "组内文本"
    inner.name = "缩小组的子形状"
    set_group_xfrm(inside, off=(0, 0), ext=(100, 100), ch_off=(0, 0), ch_ext=(200, 200))

    outside = slide.shapes.add_group_shape()              # 缩放 2：子形状换算后超出画布
    blown = outside.shapes.add_textbox(Pt(490), Pt(0), Pt(50), Pt(50))
    blown.text_frame.text = "放大后越界"
    blown.name = "放大组的子形状"
    set_group_xfrm(outside, off=(0, 0), ext=(200, 200), ch_off=(0, 0), ch_ext=(100, 100))

    result = check_presentation(save(prs, "group.pptx"))
    hits = findings(result, "fail", CHECK_BOUNDS)
    joined = " ".join(f["message"] for f in hits)
    assert "放大组的子形状" in joined, f"换算后越界的子形状必须被报出：{hits}"
    assert "缩小组的子形状" not in joined, f"换算后仍在画布内的子形状不该报：{hits}"
    assert len(hits) == 1, f"应只报 1 项越界：{hits}"
    return f"（子形状坐标换算正确：{hits[0]['message'][:52]}…）"


# --------------------------------------------------------------------------
# 用例 9：字体清单
# --------------------------------------------------------------------------
def test_font_inventory() -> str:
    prs = new_deck()
    slide = add_slide(prs)
    box = add_text(slide, 48, 24, 864, 84, "标题用雅黑", 32)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "，这段用了 Arial"
    run.font.size = Pt(24)
    run.font.name = "Arial"
    result = check_presentation(save(prs, "fonts.pptx"))
    listed = findings(result, "info", CHECK_FONTS)
    assert "Arial" in result["fonts"], f"fonts 字段应含 Arial：{result['fonts']}"
    assert any("Arial" in f["message"] for f in listed), f"非雅黑字体应被列出：{listed}"
    return f"（fonts={result['fonts']}）"


# --------------------------------------------------------------------------
# 用例 10：CLI 契约
# --------------------------------------------------------------------------
def test_cli_contract() -> str:
    qa_py = os.path.join(SCRIPTS, "qa.py")
    ok_deck = demo_deck()
    bad_prs = new_deck()
    add_text(add_slide(bad_prs), 48, 24, 864, 84, "TODO 未完成", 32)
    bad_deck = save(bad_prs, "cli_bad.pptx")
    not_pptx = os.path.join(tmp_dir(), "not_a.pptx")
    with open(not_pptx, "w", encoding="utf-8") as handle:
        handle.write("这不是 pptx\n")

    good = subprocess.run([PYTHON, qa_py, ok_deck], capture_output=True, text=True)
    assert good.returncode == 0, f"正常 deck 退出码应为 0：{good.stdout[-400:]}"
    assert "PPT 质检报告" in good.stdout, "stdout 应有可读报告"

    broken = subprocess.run([PYTHON, qa_py, bad_deck], capture_output=True, text=True)
    assert broken.returncode == 1, f"有 fail 时退出码应为 1：{broken.stdout[-400:]}"

    unreadable = subprocess.run([PYTHON, qa_py, not_pptx], capture_output=True, text=True)
    assert unreadable.returncode == 2, f"坏文件退出码应为 2：{unreadable.stderr[-300:]}"

    out_path = os.path.join(tmp_dir(), "report.json")
    streamed = subprocess.run([PYTHON, qa_py, bad_deck, "--json", "--json-out", out_path,
                               "--notes-min", "5"], capture_output=True, text=True)
    payload = json.loads(streamed.stdout)
    assert set(payload) == {"file", "slide_count", "findings", "counts", "fonts", "errors"}, \
        f"JSON 字段应与 schema 一致：{sorted(payload)}"
    assert payload["counts"]["fail"] >= 1 and payload["slide_count"] == 1
    with open(out_path, encoding="utf-8") as handle:
        assert json.load(handle)["file"] == payload["file"], "--json-out 应写出同一份结果"
    return "（退出码 0 / 1 / 2 均符合契约，JSON schema 一致）"


# --------------------------------------------------------------------------
# 运行器
# --------------------------------------------------------------------------
def main() -> int:
    """跑全部用例，返回进程退出码。"""
    cases: List[tuple] = [
        ("正常 deck 无 fail", test_normal_deck_has_no_fail),
        ("形状越界报 fail", test_out_of_bounds),
        ("缺讲者备注报 warn 并列页码", test_missing_notes),
        ("占位残留 TODO 报 fail", test_placeholder_text),
        ("空文本框 / 装饰色块口径", test_empty_textbox_versus_decoration),
        ("图片分辨率估算", test_image_dpi),
        ("标题重复 fail / 相似 warn", test_duplicate_titles),
        ("组合形状坐标换算", test_group_shape_transform),
        ("字体清单", test_font_inventory),
        ("CLI 退出码与 JSON 契约", test_cli_contract),
    ]
    failures = 0
    for name, case in cases:
        try:
            detail = case() or ""
        except Exception as exc:
            failures += 1
            print(f"FAIL  {name}")
            traceback.print_exc()
            print(f"      → {type(exc).__name__}: {exc}")
        else:
            print(f"PASS  {name} {detail}")
    print("-" * 60)
    print(f"用例 {len(cases)} 个：通过 {len(cases) - failures} 个，失败 {failures} 个")
    return 1 if failures else 0


if __name__ == "__main__":
    code = main()
    if os.environ.get("QA_KEEP_TMP") != "1" and _TMP_DIR and os.path.isdir(_TMP_DIR):
        shutil.rmtree(_TMP_DIR, ignore_errors=True)
    sys.exit(code)
